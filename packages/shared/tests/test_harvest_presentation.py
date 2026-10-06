"""Harvested slides carry an optional presentation plan without Course Studio."""

from __future__ import annotations

import ast
import json
from pathlib import Path

import pytest

from aoep_shared.harvest.export import export_course_package, presentation_speaker_notes
from aoep_shared.harvest.extractors import extract_text
from aoep_shared.harvest.generate import (
    _TEXT_DIAGRAM_STYLES,
    GeneratedCourse,
    GeneratedSlide,
    apply_harvest_presentation,
    generate_course,
    narration_sentences,
    partition_course_into_lessons,
)

_SAMPLE = (
    "Introduction\nWelcome to algebra. We cover the core objectives.\n\n"
    "Definitions\nA variable stands for an unknown number.\n\n"
    "Example 1\nA worked example solving for x.\n\n"
    "Summary\nIn summary, algebra solves for unknowns.\n"
)


def _bare(**overrides) -> GeneratedSlide:
    fields = dict(
        title="Idea",
        body="Key idea: keep it short.",
        narration="First sentence. Second sentence.",
        category="concept",
    )
    fields.update(overrides)
    return GeneratedSlide(**fields)


def _course(slides) -> GeneratedCourse:
    return GeneratedCourse(
        course_id="pres1",
        title="Presentation course",
        subject="math",
        language="en",
        source="unit",
        fmt="lecture",
        slides=slides,
    )


def test_serialization_omits_empty_presentation_fields():
    old = {"title": "T", "body": "B", "narration": "N.", "category": "concept"}
    slide = GeneratedSlide.from_dict(old)
    data = slide.to_dict()
    assert data["title"] == "T"
    assert "presentation_style_id" not in data
    assert "visual_timeline" not in data
    assert "cue_summary" not in data
    assert "checkpoints" not in data
    assert presentation_speaker_notes(slide) == []


def test_serialization_includes_style_timeline_summary_and_checkpoints():
    slide = _bare(
        presentation_style_id="process_build",
        visual_timeline={
            "source": "explicit",
            "cues": [{
                "sentence_start": 0,
                "sentence_end": 9,
                "action": "enter",
                "summary": "First sentence.",
                "start_s": -2,
                "duration_s": 1.25,
            }],
            "duration_s": 4,
        },
        checkpoints=[{"placement": "section", "activity": "try_it", "category": "exercise"}],
    )
    data = slide.to_dict()
    assert data["presentation_style_id"] == "process_build"
    timeline = data["visual_timeline"]
    assert timeline["version"] == 1
    assert timeline["source"] == "explicit"
    assert timeline["duration_s"] == 4
    assert timeline["static_frame"]["kind"] == "diagram"
    assert timeline["static_frame"]["primary_text"].startswith("Key idea:")
    cue = timeline["cues"][0]
    assert cue["sentence_start"] == 0
    assert cue["sentence_end"] == 1  # clamped to the second sentence
    assert "start_s" not in cue
    assert cue["duration_s"] == 1.25
    assert data["cue_summary"] == timeline["cue_summary"]
    assert data["cue_summary"].startswith("0-1 enter:")
    assert data["checkpoints"] == [
        {"placement": "section", "activity": "try_it", "category": "exercise"},
    ]
    again = GeneratedSlide.from_dict(data).to_dict()
    assert again["presentation_style_id"] == data["presentation_style_id"]
    assert again["visual_timeline"]["cues"] == timeline["cues"]
    assert again["checkpoints"] == data["checkpoints"]


def test_explicit_plan_is_kept_and_inferred_styles_do_not_repeat():
    authored = _bare(
        title="Quote",
        presentation_style_id="quote_card",
        checkpoints=[],
        visual_timeline={
            "source": "curated",
            "style_id": "quote_card",
            "cues": [{"sentence_start": 0, "sentence_end": 0, "action": "enter", "summary": "First sentence."}],
        },
    )
    follow = _bare(title="Next", category="concept")
    apply_harvest_presentation([authored, follow])
    assert authored.presentation_style_id == "quote_card"
    assert authored.checkpoints == []
    assert "checkpoints" not in authored.to_dict()
    assert authored.visual_timeline["source"] == "curated"
    assert follow.presentation_style_id == "sentence_highlight"
    assert follow.presentation_style_id != authored.presentation_style_id
    assert follow.visual_timeline["static_frame"]["kind"] == "text"
    assert follow.checkpoints == [
        {"placement": "lesson", "activity": "close", "category": "concept"},
    ]

    first = _bare(title="A", category="concept")
    second = _bare(title="B", category="concept")
    apply_harvest_presentation([first, second])
    assert first.presentation_style_id == "sentence_highlight"
    assert second.presentation_style_id != first.presentation_style_id
    assert second.presentation_style_id in _TEXT_DIAGRAM_STYLES


def test_generate_course_adds_text_diagram_plan_and_checkpoints():
    course = generate_course(extract_text(_SAMPLE, default_title="Algebra 101"), subject="math")
    styles = [slide.presentation_style_id for slide in course.slides]
    assert styles
    assert all(style in _TEXT_DIAGRAM_STYLES for style in styles)
    assert all(left != right for left, right in zip(styles, styles[1:]))
    assert course.slides[0].category == "introduction"
    assert "checkpoints" not in course.slides[0].to_dict()
    exercises = [slide for slide in course.slides if slide.category == "exercise"]
    assert exercises
    for slide in exercises:
        markers = slide.to_dict()["checkpoints"]
        assert any(row["placement"] == "section" and row["activity"] == "try_it" for row in markers)
    closing = course.slides[-1].to_dict()
    assert any(row["placement"] == "lesson" and row["activity"] == "close" for row in closing["checkpoints"])
    for slide in course.slides:
        payload = slide.to_dict()
        count = len(narration_sentences(slide.narration)) or 1
        timeline = payload["visual_timeline"]
        assert timeline["version"] == 1
        assert timeline["static_frame"]["kind"] in {"text", "diagram"}
        assert payload["cue_summary"]
        for cue in timeline["cues"]:
            assert 0 <= cue["sentence_start"] <= cue["sentence_end"] < count


def test_media_slide_uses_image_style_and_media_static_frame():
    slide = _bare(category="concept", media_url="https://example.com/demo.mp4", media_kind="video")
    apply_harvest_presentation([slide])
    data = slide.to_dict()
    assert data["presentation_style_id"] == "annotated_image"
    assert data["visual_timeline"]["static_frame"]["kind"] == "media"
    assert data["visual_timeline"]["static_frame"]["media_url"].endswith("demo.mp4")


def test_partition_keeps_slide_presentation_metadata():
    marked = _bare(
        title="Marked section",
        category="exercise",
        presentation_style_id="sentence_highlight",
        visual_timeline={
            "source": "explicit",
            "style_id": "sentence_highlight",
            "cues": [{"sentence_start": 0, "sentence_end": 0, "summary": "First sentence."}],
        },
        checkpoints=[{"placement": "section", "activity": "try_it", "category": "exercise"}],
    )
    slides = [
        GeneratedSlide(title=f"Section {i}", body="Body.", narration=f"Narration {i}.", category="concept")
        for i in range(29)
    ]
    slides.insert(3, marked)
    lessons = partition_course_into_lessons(_course(slides), max_slides=12)
    assert len(lessons) == 3
    assert sum(len(lesson.slides) for lesson in lessons) == 30
    found = [slide for lesson in lessons for slide in lesson.slides if slide.title == "Marked section"]
    assert found == [marked]
    payload = found[0].to_dict()
    assert payload["presentation_style_id"] == "sentence_highlight"
    assert payload["visual_timeline"]["version"] == 1
    assert payload["cue_summary"]
    assert payload["checkpoints"][0]["activity"] == "try_it"


def test_speaker_note_blocks_follow_legacy_notes():
    bare = _bare()
    assert presentation_speaker_notes(bare) == []

    planned = _bare(
        category="summary",
        presentation_style_id="recap_layout",
        body="Takeaway: practice once.",
        narration="You made it. Practice today.",
        audio_path="audio/close.aiff",
        checkpoints=[{"placement": "lesson", "activity": "close", "category": "summary"}],
        visual_timeline={
            "source": "inferred",
            "style_id": "recap_layout",
            "cues": [
                {"sentence_start": 0, "sentence_end": 0, "action": "enter", "summary": "You made it."},
                {"sentence_start": 1, "sentence_end": 1, "action": "emphasis", "summary": "Practice today."},
            ],
        },
    )
    blocks = presentation_speaker_notes(planned)
    assert blocks[0].splitlines()[0] == "[Presentation style]"
    assert "id: recap_layout" in blocks[0]
    assert "static_frame: text" in blocks[0]
    assert "primary: Takeaway: practice once." in blocks[0]
    assert "checkpoint: lesson:close" in blocks[0]
    assert blocks[1].splitlines()[0] == "[Visual timeline v1]"
    assert "cue_summary:" in blocks[1]
    assert "0 enter text | You made it." in blocks[1]
    assert "1 emphasis text | Practice today." in blocks[1]


def test_pptx_notes_keep_legacy_text_and_append_plan_blocks(tmp_path):
    pytest.importorskip("pptx")
    from pptx import Presentation

    legacy = _bare(title="Legacy", narration="Say this plainly.", category="concept")
    planned = _bare(
        title="Planned",
        narration="Watch the diagram. Then name the step.",
        category="demo",
        presentation_style_id="diagram_build",
        audio_path="voices/demo.aiff",
        media_url="https://example.com/walk.mp4",
        media_kind="video",
        checkpoints=[{"placement": "section", "activity": "try_it", "category": "exercise"}],
        visual_timeline={
            "source": "inferred",
            "cues": [{
                "sentence_start": 0,
                "sentence_end": 1,
                "action": "enter",
                "layer": "diagram",
                "summary": "Watch the diagram.",
            }],
        },
    )
    # Body long enough to keep the legacy full-reference note in front of the plan.
    planned.body = "Step 1\n" + ("detail " * 80)
    pkg = export_course_package(_course([legacy, planned]), tmp_path, write_pptx=True)
    prs = Presentation(str(pkg.pptx_path))
    content = list(prs.slides)[1:]
    legacy_notes = content[0].notes_slide.notes_text_frame.text
    planned_notes = content[1].notes_slide.notes_text_frame.text

    assert legacy_notes == "Say this plainly.\n\n[Slide type: concept]"
    assert "[Presentation style]" not in legacy_notes
    assert "[Visual timeline" not in legacy_notes

    assert planned_notes.index("Watch the diagram.") < planned_notes.index("[Audio narration:")
    assert planned_notes.index("[Audio narration:") < planned_notes.index("[Video demo:")
    assert planned_notes.index("[Video demo:") < planned_notes.index("[Slide type: demo]")
    assert planned_notes.index("[Slide type: demo]") < planned_notes.index("[Full reference text]")
    assert planned_notes.index("[Full reference text]") < planned_notes.index("[Presentation style]")
    assert planned_notes.index("[Presentation style]") < planned_notes.index("[Visual timeline v1]")
    assert "static_frame: media" in planned_notes
    assert "media: https://example.com/walk.mp4" in planned_notes
    assert "checkpoint: section:try_it" in planned_notes
    assert "id: diagram_build" in planned_notes
    # The on-slide frame stays the static title and bullets.
    assert content[1].shapes.title.text == "Planned"
    shape_text = "\n".join(shape.text_frame.text for shape in content[1].shapes if shape.has_text_frame)
    assert "[Visual timeline" not in shape_text


def test_course_json_exports_presentation_metadata(tmp_path):
    course = generate_course(extract_text(_SAMPLE, default_title="Algebra 101"), subject="math")
    pkg = export_course_package(course, tmp_path, write_pptx=False)
    data = json.loads(pkg.course_json_path.read_text())
    first = data["slides"][0]
    assert first["presentation_style_id"]
    assert first["visual_timeline"]["version"] == 1
    assert first["visual_timeline"]["static_frame"]["kind"] in {"text", "diagram"}
    assert first["cue_summary"]
    assert any(slide.get("checkpoints") for slide in data["slides"])


def test_harvest_modules_do_not_import_course_studio():
    root = Path(__file__).resolve().parents[1] / "src" / "aoep_shared" / "harvest"
    for name in ("generate.py", "export.py"):
        tree = ast.parse((root / name).read_text())
        modules = []
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                modules.extend(alias.name for alias in node.names)
            elif isinstance(node, ast.ImportFrom):
                modules.append(node.module or "")
        assert all("theodore_course_studio" not in module for module in modules)
